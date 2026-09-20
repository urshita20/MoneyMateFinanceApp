import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def build_moneymate_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6]

    # Colors
    BG_COLOR = RGBColor(9, 13, 22)         # #090D16
    CARD_BG = RGBColor(17, 24, 39)         # #111827
    CARD_BORDER = RGBColor(30, 41, 59)     # #1E293B
    EMERALD = RGBColor(16, 185, 129)       # #10B981
    WHITE = RGBColor(255, 255, 255)        # #FFFFFF
    MUTED = RGBColor(148, 163, 184)        # #94A3B8
    LIGHT_GRAY = RGBColor(226, 232, 240)   # #E2E8F0

    def set_slide_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = BG_COLOR

    def add_header(slide, tag_text, title_text, subtitle_text):
        # Tag
        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.4))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = tag_text.upper()
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = EMERALD

        # Title
        txBox2 = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.8))
        tf2 = txBox2.text_frame
        tf2.word_wrap = True
        p2 = tf2.paragraphs[0]
        p2.text = title_text
        p2.font.size = Pt(28)
        p2.font.bold = True
        p2.font.color.rgb = WHITE

        # Subtitle
        txBox3 = slide.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.7), Inches(0.6))
        tf3 = txBox3.text_frame
        tf3.word_wrap = True
        p3 = tf3.paragraphs[0]
        p3.text = subtitle_text
        p3.font.size = Pt(14)
        p3.font.color.rgb = MUTED

    def add_card(slide, left, top, width, height, icon_str, title_str, desc_str):
        # Card Background Shape
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = CARD_BG
        shape.line.color.rgb = CARD_BORDER
        shape.line.width = Pt(1)

        # Card Text Frame
        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.25)
        tf.margin_bottom = Inches(0.25)

        # Icon / Title
        p0 = tf.paragraphs[0]
        p0.text = f"{icon_str}  {title_str}"
        p0.font.size = Pt(16)
        p0.font.bold = True
        p0.font.color.rgb = WHITE
        p0.space_after = Pt(10)

        # Desc
        p1 = tf.add_paragraph()
        p1.text = desc_str
        p1.font.size = Pt(12)
        p1.font.color.rgb = MUTED

    # ==========================================
    # SLIDE 1: Title & Cover Slide
    # ==========================================
    slide1 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide1)

    txBox = slide1.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(11.7), Inches(0.5))
    p = txBox.text_frame.paragraphs[0]
    p.text = "EXECUTIVE PRODUCT PRESENTATION"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = EMERALD
    p.alignment = PP_ALIGN.CENTER

    txBoxTitle = slide1.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(1.2))
    pTitle = txBoxTitle.text_frame.paragraphs[0]
    pTitle.text = "MoneyMate"
    pTitle.font.size = Pt(54)
    pTitle.font.bold = True
    pTitle.font.color.rgb = WHITE
    pTitle.alignment = PP_ALIGN.CENTER

    txBoxSub = slide1.shapes.add_textbox(Inches(1.5), Inches(3.0), Inches(10.3), Inches(1.0))
    pSub = txBoxSub.text_frame.paragraphs[0]
    pSub.text = "The Next-Generation Personal Wealth & AI Financial Intelligence Platform\nBuilt for Real Data Clarity, Automated Tracking, Household Security & Youth Financial Literacy"
    pSub.font.size = Pt(16)
    pSub.font.color.rgb = MUTED
    pSub.alignment = PP_ALIGN.CENTER

    # 3 Stat Cards
    add_card(slide1, Inches(1.0), Inches(4.5), Inches(3.5), Inches(2.0), "🤖", "AI Copilot & Insights", "Personalized financial advisory powered by real user analytics & smart budgeting logic.")
    add_card(slide1, Inches(4.9), Inches(4.5), Inches(3.5), Inches(2.0), "📊", "Full Wealth Suite", "Statement OCR, recurring commitment detector, safety net planner & time machine simulator.")
    add_card(slide1, Inches(8.8), Inches(4.5), Inches(3.5), Inches(2.0), "🐣", "MoneyMate Junior", "Gamified child financial literacy mode protected by parental key security.")


    # ==========================================
    # SLIDE 2: Financial Dashboard
    # ==========================================
    slide2 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide2)
    add_header(slide2, "Module 01", "Unified Financial Dashboard", "A real-time command center giving users complete control over cash flow, net worth, and high-frequency quick actions.")

    add_card(slide2, Inches(0.8), Inches(2.4), Inches(5.6), Inches(2.1), "📈", "Real-Time Income vs. Expense", "Instant visual feedback on total monthly earnings, category expenditure, and live savings rate calculations with synchronized health metrics.")
    add_card(slide2, Inches(6.8), Inches(2.4), Inches(5.6), Inches(2.1), "⚡", "One-Tap Quick Actions", "Direct shortcuts to add income/expense, launch the AI Copilot, run the Wealth Time Machine, or check your Financial Health Score.")
    add_card(slide2, Inches(0.8), Inches(4.8), Inches(5.6), Inches(2.1), "🎨", "Modern Glassmorphism UI", "Vibrant dark mode UI built with modern typography, smooth micro-interactions, and responsive layout standards.")
    add_card(slide2, Inches(6.8), Inches(4.8), Inches(5.6), Inches(2.1), "🛡️", "Privacy-First Architecture", "End-to-end data security ensuring no sensitive financial records are shared or stored unencrypted.")


    # ==========================================
    # SLIDE 3: OCR & Statement Scanner
    # ==========================================
    slide3 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide3)
    add_header(slide3, "Module 02", "Smart Import & OCR Scanner", "Eliminate manual data entry with automated bank statement parsing and camera OCR receipt scanning.")

    add_card(slide3, Inches(0.8), Inches(2.4), Inches(5.6), Inches(2.1), "📄", "Automated Bank Statement Import", "Upload multi-bank PDF/CSV/Excel statements. Automatically extracts transactions, dates, amounts, and merchant tags seamlessly.")
    add_card(slide3, Inches(6.8), Inches(2.4), Inches(5.6), Inches(2.1), "📷", "Intelligent Receipt OCR", "Scan physical receipts using device camera or file upload. Extracts merchant name, total price, line items, and date instantly.")
    add_card(slide3, Inches(0.8), Inches(4.8), Inches(5.6), Inches(2.1), "🏷️", "Auto-Categorization Engine", "Intelligently classifies transactions into Food, Utilities, Shopping, Bills, Transport, and Entertainment without manual effort.")
    add_card(slide3, Inches(6.8), Inches(4.8), Inches(5.6), Inches(2.1), "🔍", "Duplicate & Anomaly Shield", "Prevents double-counting imported statement entries and flags suspicious or unusual spending spikes automatically.")


    # ==========================================
    # SLIDE 4: Bills & Commitments
    # ==========================================
    slide4 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide4)
    add_header(slide4, "Module 03", "Bills & Commitments Radar", "Never miss a bill or EMI payment again. Automated detection of recurring subscriptions, utilities, and financial commitments.")

    add_card(slide4, Inches(0.8), Inches(2.8), Inches(3.6), Inches(3.8), "🔄", "Auto-Subscription Detection", "Smart algorithms scan transaction history to auto-identify Netflix, Spotify, Gym, SIPs, EMIs, and recurring bills automatically.")
    add_card(slide4, Inches(4.8), Inches(2.8), Inches(3.6), Inches(3.8), "📅", "Upcoming Due Calendar", "Organized chronological list of due dates with status indicators (Paid, Upcoming Dues, Pending Payments) to keep bills on track.")
    add_card(slide4, Inches(8.8), Inches(2.8), Inches(3.6), Inches(3.8), "🚨", "Cash Flow Defense", "Proactively warns users if upcoming commitments exceed predicted account balance to prevent overdraft fees and penalties.")


    # ==========================================
    # SLIDE 5: Health Score & Safety Net
    # ==========================================
    slide5 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide5)
    add_header(slide5, "Module 04", "Financial Health & Safety Net", "Comprehensive financial diagnostic score paired with a data-driven survival planner for emergencies.")

    add_card(slide5, Inches(0.8), Inches(2.4), Inches(5.6), Inches(4.5), "🩺", "Algorithmic Health Score (0–100)", "Evaluates 4 critical financial pillars:\n\n• Savings Rate Performance\n• Emergency Fund Cushion\n• Debt-to-Income Ratio\n• Budget Adherence Discipline\n\nProvides real-time actionable steps to continuously improve financial fitness scores.")
    add_card(slide5, Inches(6.8), Inches(2.4), Inches(5.6), Inches(4.5), "🛟", "Financial Survival Planner", "Calculates exact monthly survival expense baselines:\n\n• Essential Housing & Food Baseline\n• Emergency Fund Target Calculator\n• Cash Runway Months Estimation\n• Burn Rate Optimization\n\nProtects households against unexpected income shocks and emergencies.")


    # ==========================================
    # SLIDE 6: Time Machine Simulator
    # ==========================================
    slide6 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide6)
    add_header(slide6, "Module 05", "Wealth Time Machine Simulator", "Interactive future wealth forecasting engine. Project net worth expansion over 5, 10, 20, or 30 years.")

    add_card(slide6, Inches(0.8), Inches(2.4), Inches(5.6), Inches(4.5), "⏳", "Compound Growth Projections", "Simulates long-term wealth trajectory:\n\n• Monthly Savings & SIP Compounding\n• Expected Yield & Annual Inflation Adjustments\n• 5, 10, 20, 30 Year Interactive Timelines\n\nAllows users to visualize exact net worth milestones before making major investments.")
    add_card(slide6, Inches(6.8), Inches(2.4), Inches(5.6), Inches(4.5), "🎯", "Life Milestone Modeler", "Tests real-world financial decision scenarios:\n\n• 'What if I buy a home in 3 years?'\n• 'What if I increase monthly SIP by 15%?'\n• 'What if I take a 6-month career sabbatical?'\n\nProvides instant risk-adjusted outcome forecasts.")


    # ==========================================
    # SLIDE 7: MoneyMate Junior
    # ==========================================
    slide7 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide7)
    add_header(slide7, "Module 06", "MoneyMate Junior — Youth Financial Literacy", "A dedicated, kid-friendly financial literacy experience empowering children to learn smart saving habits early.")

    add_card(slide7, Inches(0.8), Inches(2.8), Inches(3.6), Inches(3.8), "🔐", "Parental PIN Security", "Parental Key gatekeeping ensures only parents can deposit allowance payouts, adjust spending limits, or manage settings.")
    add_card(slide7, Inches(4.8), Inches(2.8), Inches(3.6), Inches(3.8), "🐖", "Allowance & Goal Saver", "Visual progress bars for kids' savings goals (toys, bikes, games, books) with celebratory micro-animations on milestones.")
    add_card(slide7, Inches(8.8), Inches(2.8), Inches(3.6), Inches(3.8), "⭐", "Chore Rewards & Quests", "Gamified task completion engine rewarding daily responsibility and financial discipline with allowance payouts.")


    # ==========================================
    # SLIDE 8: Knowledge Hub & Live Deployment
    # ==========================================
    slide8 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide8)
    add_header(slide8, "Module 07 & Conclusion", "Knowledge Hub & Live Deployment", "Tailored financial education, tax regime calculators, and direct access to the live cloud application.")

    add_card(slide8, Inches(0.8), Inches(2.4), Inches(5.6), Inches(2.3), "📚", "Personalized Knowledge Hub", "Curated articles, tax guides comparing New vs. Old Tax Regimes, debt payoff strategies, and beginner SIP investment rules.")
    add_card(slide8, Inches(6.8), Inches(2.4), Inches(5.6), Inches(2.3), "🚀", "Live Vercel Cloud Deployment", "Fully built, optimized, and deployed live on Vercel with automated CI/CD and sub-second loading performance.")

    # Call to action card
    shape = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.0), Inches(11.6), Inches(1.8))
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD_BG
    shape.line.color.rgb = EMERALD
    shape.line.width = Pt(2)
    tf = shape.text_frame
    tf.word_wrap = True
    p0 = tf.paragraphs[0]
    p0.text = "🔗 Experience MoneyMate Live"
    p0.font.size = Pt(20)
    p0.font.bold = True
    p0.font.color.rgb = WHITE
    p0.alignment = PP_ALIGN.CENTER
    p0.space_after = Pt(8)

    p1 = tf.add_paragraph()
    p1.text = "https://money-mate-finance-app.vercel.app/"
    p1.font.size = Pt(16)
    p1.font.bold = True
    p1.font.color.rgb = EMERALD
    p1.alignment = PP_ALIGN.CENTER

    output_path = os.path.join("public", "MoneyMate_Presentation.pptx")
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    build_moneymate_presentation()
